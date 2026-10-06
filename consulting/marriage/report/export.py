"""PDF and DOCX export for a stored marriage consultation."""

from __future__ import annotations

import html
import re
import shutil
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Literal, Protocol

from docx import Document
from docx.shared import Cm, Pt

from consulting.marriage.dto.history import MarriageStoredResult
from consulting.marriage.models.report import ReportSection
from consulting.marriage.report.access import customer_sections
from engines.report_engine.contracts.report_export_result_v1 import MEDIA_TYPE_DOCX, MEDIA_TYPE_PDF
from engines.report_engine.exporting.docx_exporter_v1 import validate_docx_file
from engines.report_engine.exporting.pdf_exporter_v1 import (
    PlaywrightPdfBackend,
    validate_pdf_file,
)

MarriageExportFormat = Literal["pdf", "docx"]


class MarriagePdfBackend(Protocol):
    """Small PDF boundary used by the marriage document exporter."""

    def html_to_pdf(self, content: str, output_path: Path, *, title: str) -> int | None:
        """Write the rendered PDF and optionally return its page count."""


@dataclass(frozen=True, slots=True)
class MarriageExportArtifact:
    """Temporary downloadable file metadata."""

    path: Path
    filename: str
    media_type: str


@dataclass(frozen=True, slots=True)
class ConsultationExportDocument:
    title: str
    heading: str
    subtitle: str
    filename_stem: str
    summary_rows: list[tuple[str, str]]
    sections: list[ReportSection]
    notice: str


def export_marriage_file(
    stored: MarriageStoredResult,
    fmt: MarriageExportFormat,
    *,
    output_root: Path | None = None,
    pdf_backend: MarriagePdfBackend | None = None,
    document: ConsultationExportDocument | None = None,
) -> MarriageExportArtifact:
    """Render the stored customer report without recomputing the consultation."""
    if stored.report_model is None:
        raise ValueError("marriage_report_missing")
    document = document or ConsultationExportDocument(
        title=_title(stored),
        heading="BẢN LUẬN GIẢI TƯ VẤN HÔN NHÂN",
        subtitle=_title(stored).replace("Tư vấn hôn nhân - ", ""),
        filename_stem=_filename(stored, fmt).rsplit(".", 1)[0],
        summary_rows=_summary_rows(stored),
        sections=_visible_sections(stored),
        notice="Kết quả phản ánh mô hình luận giải Bát Tự và Cung Phi; không phải xác suất hạnh phúc hay quyết định thay cho hai người.",
    )
    return export_consultation_file(
        document, stored.history.consultation_id, fmt,
        output_root=output_root, pdf_backend=pdf_backend,
    )


def export_consultation_file(
    document: ConsultationExportDocument,
    consultation_id: str,
    fmt: MarriageExportFormat,
    *,
    output_root: Path | None = None,
    pdf_backend: MarriagePdfBackend | None = None,
) -> MarriageExportArtifact:
    """Render the common consultation layout from an archived document."""
    root = output_root or Path(tempfile.mkdtemp(prefix="bte-marriage-export-"))
    root.mkdir(parents=True, exist_ok=True)
    filename = f"{document.filename_stem}.{fmt}"
    path = root / filename
    try:
        if fmt == "docx":
            _write_docx(document, consultation_id, path)
            validate_docx_file(path)
            media_type = MEDIA_TYPE_DOCX
        elif fmt == "pdf":
            backend = pdf_backend or PlaywrightPdfBackend()
            backend.html_to_pdf(_render_html(document), path, title=document.title)
            validate_pdf_file(path)
            media_type = MEDIA_TYPE_PDF
        else:
            raise ValueError(f"unsupported_marriage_export:{fmt}")
    except Exception:
        if output_root is None:
            cleanup_marriage_export(path)
        raise
    return MarriageExportArtifact(path=path, filename=filename, media_type=media_type)


def cleanup_marriage_export(path: Path) -> None:
    """Remove one temporary export directory after the response completes."""
    try:
        shutil.rmtree(path.parent, ignore_errors=True)
    except OSError:
        return


def _title(stored: MarriageStoredResult) -> str:
    left = stored.result.person_a.display_name or "Người Nữ"
    right = stored.result.person_b.display_name or "Người Nam"
    return f"Tư vấn hôn nhân - {left} và {right}"


def _filename(stored: MarriageStoredResult, fmt: MarriageExportFormat) -> str:
    left = stored.result.person_a.display_name or "nguoi-nu"
    right = stored.result.person_b.display_name or "nguoi-nam"
    identity = _slug(f"{left}-{right}") or "cap-doi"
    return f"BTE_TuVanHonNhan_{identity}.{fmt}"


def _slug(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Za-z0-9]+", "-", normalized).strip("-").lower()


def _visible_sections(stored: MarriageStoredResult) -> list[ReportSection]:
    model = stored.report_model
    if model is None:
        return []
    return [item for item in customer_sections(model) if item.section_id != "appendix"]


def _write_docx(spec: ConsultationExportDocument, consultation_id: str, path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.core_properties.title = spec.title
    document.core_properties.subject = spec.heading
    document.core_properties.identifier = consultation_id

    document.add_heading(spec.heading, level=0)
    document.add_paragraph(spec.subtitle)
    summary = spec.summary_rows
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in summary:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    for report_section in spec.sections:
        document.add_heading(report_section.title, level=1)
        if report_section.summary:
            document.add_paragraph(report_section.summary)
        for block in report_section.blocks:
            if block.visibility != "customer" or block.kind == "methodology":
                continue
            if block.title and block.title != block.body:
                document.add_heading(block.title, level=2)
            if block.body:
                document.add_paragraph(block.body)
    document.add_paragraph()
    document.add_paragraph(
        spec.notice
    )
    document.save(path)


def _summary_rows(stored: MarriageStoredResult) -> list[tuple[str, str]]:
    score = "Chưa chấm điểm" if stored.history.score is None else f"{stored.history.score:g}/100"
    grade = stored.history.grade.value if stored.history.grade else "—"
    return [
        ("Mã hồ sơ", stored.history.consultation_id),
        ("Ngày lập", stored.history.created_at),
        ("Điểm tương hợp", score),
        ("Xếp loại", grade),
    ]


def _render_html(spec: ConsultationExportDocument) -> str:
    section_html: list[str] = []
    for report_section in spec.sections:
        blocks: list[str] = []
        if report_section.summary:
            blocks.append(f"<p>{html.escape(report_section.summary)}</p>")
        for block in report_section.blocks:
            if block.visibility != "customer" or block.kind == "methodology":
                continue
            heading = ""
            if block.title and block.title != block.body:
                heading = f"<h3>{html.escape(block.title)}</h3>"
            body = f"<p>{html.escape(block.body).replace(chr(10), '<br>')}</p>" if block.body else ""
            blocks.append(f"{heading}{body}")
        section_html.append(
            f"<section><h2>{html.escape(report_section.title)}</h2>{''.join(blocks)}</section>"
        )
    rows = "".join(
        f"<tr><th>{html.escape(label)}</th><td>{html.escape(value)}</td></tr>"
        for label, value in spec.summary_rows
    )
    return f"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8"><title>{html.escape(spec.title)}</title>
<style>
@page {{ size: A4; margin: 18mm; }}
body {{ font-family: Arial, 'Segoe UI', sans-serif; color: #172033; font-size: 11pt; line-height: 1.55; }}
h1 {{ font-size: 22pt; margin: 0 0 8px; }} h2 {{ font-size: 15pt; margin: 22px 0 7px; border-bottom: 1px solid #d9dee8; padding-bottom: 5px; }}
h3 {{ font-size: 12pt; margin: 13px 0 4px; }} p {{ margin: 5px 0 9px; }}
h1, h2, h3 {{ break-after: avoid; }} p {{ orphans: 3; widows: 3; overflow-wrap: anywhere; }} tr {{ break-inside: avoid; }}
table {{ width: 100%; border-collapse: collapse; margin: 16px 0 22px; }} th, td {{ border: 1px solid #d9dee8; padding: 7px 9px; text-align: left; }} th {{ width: 30%; background: #f4f7fb; }}
.subtitle {{ color: #526078; font-size: 13pt; }} .notice {{ margin-top: 24px; color: #526078; font-size: 9.5pt; }}
</style></head><body>
<h1>{html.escape(spec.heading)}</h1><p class="subtitle">{html.escape(spec.subtitle)}</p>
<table>{rows}</table>{''.join(section_html)}
<p class="notice">{html.escape(spec.notice)}</p>
</body></html>"""
