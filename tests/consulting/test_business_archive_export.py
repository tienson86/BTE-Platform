"""Business profiles retain occupation, evidence and downloadable reports."""

import io
import json
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from docx import Document
from tests.consulting.api_fixtures import valid_body
from tests.consulting.test_marriage_archive_export import _client
from consulting.business.export import export_business_file
from consulting.marriage.api.service import MarriageConsultationApi

BASE = "/api/v1/consulting/marriage"


def _create(client, occupation="technology", label="Công nghệ - phần mềm"):
    body = valid_body(gender_a="male", gender_b="male")
    body["options"] = {"include_score": True}
    body["request_meta"] = {"client": f"business_consulting;occupation={occupation};occupation_label={label}"}
    response = client.post(BASE, json=body)
    assert response.status_code == 201
    return response.json()["data"]["consultation_id"]


def test_business_save_is_idempotent_and_survives_restart(tmp_path: Path):
    path = tmp_path / "records.json"
    client, _ = _client(path)
    id = _create(client)
    assert client.get(f"{BASE}/business/history").json()["data"]["items"] == []
    first = client.post(f"{BASE}/business/{id}/save")
    assert first.status_code == 200
    saved = first.json()["data"]
    assert saved["saved_at"] and saved["occupation"] == "technology"
    assert client.post(f"{BASE}/business/{id}/save").json()["data"]["saved_at"] == saved["saved_at"]
    reloaded, repository = _client(path)
    with patch.object(MarriageConsultationApi, "create_consultation", side_effect=AssertionError("must not recalculate")):
        response = reloaded.get(f"{BASE}/business/{id}")
        assert response.status_code == 200
        data = response.json()["data"]
        assert data["business_context"]["occupation_label"] == "Công nghệ - phần mềm"
        assert data["input"]["person_a"]["gender"] == "male"
        assert data["input"]["person_a"]["birth_date"] == "1987-01-21"
        assert len(data["assessment_cards"]) == 5
        assert data["compatibility_matrix"]["sections"]
        assert repository.get(id).business_profile.report_model.sections
        assert len(reloaded.get(f"{BASE}/business/history").json()["data"]["items"]) == 1


def test_history_filters_unsaved_and_marriage_records_before_pagination(tmp_path: Path):
    client, _ = _client(tmp_path / "records.json")
    client.post(BASE, json=valid_body())
    _create(client, "retail", "Bán lẻ")
    first = _create(client)
    second = _create(client, "manufacturing", "Sản xuất")
    for id in (first, second):
        client.post(f"{BASE}/business/{id}/save")
    page = client.get(f"{BASE}/business/history?limit=1").json()["data"]
    assert page["items"][0]["consultation_id"] == second
    assert page["next_cursor"] == second
    next_page = client.get(f"{BASE}/business/history?limit=1&cursor={second}").json()["data"]
    assert next_page["items"][0]["consultation_id"] == first
    assert next_page["next_cursor"] is None
    assert client.get(f"{BASE}/business/history?limit=0").status_code == 400


def test_docx_matches_business_report_with_occupation_and_all_evidence(tmp_path: Path):
    client, repository = _client(tmp_path / "records.json")
    id = _create(client)
    client.post(f"{BASE}/business/{id}/save")
    response = client.get(f"{BASE}/business/{id}/export/docx")
    assert response.status_code == 200
    assert "BTE_TuVanHopTac" in response.headers["content-disposition"]
    document = Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "TƯ VẤN HỢP TÁC" in text and "TƯ VẤN HÔN NHÂN" not in text
    assert "Công nghệ - phần mềm" in text
    for term in ("Cung Phi", "Dụng", "Can Chi", "Nhật Chủ", "Mệnh Cục", "Cơ sở"):
        assert term in text
    stored = repository.get(id)
    for section in stored.business_profile.report_model.sections:
        for block in section.blocks:
            assert not block.body or block.body in text


def test_saved_wording_is_used_for_display_and_export(tmp_path: Path):
    client, repository = _client(tmp_path / "records.json")
    id = _create(client)
    client.post(f"{BASE}/business/{id}/save")
    with patch("consulting.business.presentation.build_business_profile", side_effect=AssertionError("must use snapshot")):
        assert client.get(f"{BASE}/business/{id}/report").status_code == 200
        assert client.get(f"{BASE}/business/{id}").status_code == 200
    stored = repository.get(id)
    with patch("consulting.business.presentation.build_business_profile", side_effect=AssertionError("must use snapshot")):
        artifact = export_business_file(stored, "docx", output_root=tmp_path / "export")
        assert artifact.path.is_file()


def test_business_score_is_the_same_on_display_report_export_and_restart(tmp_path: Path):
    path = tmp_path / "records.json"
    client, repository = _client(path)
    id = _create(client)
    before = client.get(f"{BASE}/business/{id}/report").json()["data"]["business_score"]
    assert before["model_version"] == "business.compatibility.v1"
    assert before["score"] is None or 0 <= before["score"] <= 100
    client.post(f"{BASE}/business/{id}/save")
    client, _ = _client(path)
    profile = client.get(f"{BASE}/business/{id}").json()["data"]
    report = client.get(f"{BASE}/business/{id}/report").json()["data"]
    assert profile["business_score"] == report["business_score"] == before
    assert profile["score"] is None and report["score"] is None
    document = Document(io.BytesIO(client.get(f"{BASE}/business/{id}/export/docx").content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert before["recommendation"] in text
    assert before["methodology"] in text
    assert before["disclaimer"] in text
    assert repository.get(id).business_profile.business_score is not None


def test_legacy_saved_profile_gains_score_from_stored_evidence_without_birth_recalculation(tmp_path: Path):
    path = tmp_path / "records.json"
    client, repository = _client(path)
    id = _create(client)
    client.post(f"{BASE}/business/{id}/save")
    stored = repository.get(id)
    repository.save(replace(stored, business_profile=replace(stored.business_profile, business_score=None)))
    data = json.loads(path.read_text(encoding="utf-8"))
    # The old disk format did not contain this optional field at all.
    def strip_score(value):
        if isinstance(value, dict):
            value.pop("business_score", None)
            for child in value.values():
                strip_score(child)
        elif isinstance(value, list):
            for child in value:
                strip_score(child)
    strip_score(data)
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    reloaded, _ = _client(path)
    with patch.object(MarriageConsultationApi, "create_consultation", side_effect=AssertionError("must not recalculate")):
        profile = reloaded.get(f"{BASE}/business/{id}").json()["data"]
        report = reloaded.get(f"{BASE}/business/{id}/report").json()["data"]
    assert profile["business_context"]["saved_at"] == stored.business_profile.saved_at
    assert profile["business_score"] == report["business_score"]
    assert any(section["section_id"] == "business_score" for section in report["sections"])


def test_unrelated_profiles_and_missing_ids_are_rejected(tmp_path: Path):
    client, _ = _client(tmp_path / "records.json")
    id = client.post(BASE, json=valid_body()).json()["data"]["consultation_id"]
    for suffix in ("", "/report", "/export/docx"):
        assert client.get(f"{BASE}/business/{id}{suffix}").status_code == 404
    assert client.post(f"{BASE}/business/{id}/save").status_code == 404
    assert client.get(f"{BASE}/business/missing/export/pdf").status_code == 404


def test_save_failure_does_not_claim_success_or_leave_a_saved_memory_record(tmp_path: Path):
    client, repository = _client(tmp_path / "records.json")
    id = _create(client)
    with patch.object(repository, "_flush", side_effect=OSError("disk unavailable")):
        assert client.post(f"{BASE}/business/{id}/save").status_code == 500
    assert repository.get(id).business_profile is None
    assert client.get(f"{BASE}/business/history").json()["data"]["items"] == []
    assert client.post(f"{BASE}/business/{id}/save").status_code == 200
