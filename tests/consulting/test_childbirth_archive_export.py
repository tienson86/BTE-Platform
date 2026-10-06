"""Childbirth archives and exports retain the exact analyzed result."""

import io
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from consulting.childbirth.api import build_childbirth_router
from consulting.childbirth.export import export_childbirth_file
from consulting.childbirth.repository import JsonChildbirthRepository
from consulting.childbirth.service import analyze_childbirth_plan
from consulting.marriage.adapters.canonical_runtime import CanonicalOrchestratorAdapter
from tests.consulting.runtime_fixtures import FakeCanonicalRunner

BASE = "/api/v1/consulting/childbirth"
BODY = {
    "father": {"gender": "male", "full_name": "An", "birth_date": "1990-01-02", "birth_time": "02:00", "birth_place": "Hà Nội"},
    "mother": {"gender": "female", "full_name": "Bình", "birth_date": "1992-02-01"},
    "options": {"start_year": 2026, "years_ahead": 4},
}


def _client(directory: Path):
    repository = JsonChildbirthRepository(directory)
    app = FastAPI()
    app.include_router(build_childbirth_router(repository), prefix="/api/v1")
    return TestClient(app), repository


def _create(client):
    def analyze(**kwargs):
        return analyze_childbirth_plan(
            **kwargs, adapter=CanonicalOrchestratorAdapter(runner=FakeCanonicalRunner())
        )

    with patch("consulting.childbirth.api.analyze_childbirth_plan", side_effect=analyze):
        response = client.post(f"{BASE}/analyze", json=BODY)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_save_is_idempotent_and_profiles_survive_restart(tmp_path):
    client, _ = _client(tmp_path)
    data = _create(client)
    id = data["consultation_id"]
    assert client.get(f"{BASE}/history").json()["data"]["items"] == []
    first = client.post(f"{BASE}/{id}/save").json()["data"]
    assert first["saved_at"]
    assert client.post(f"{BASE}/{id}/save").json()["data"] == first
    reloaded, _ = _client(tmp_path)
    with patch("consulting.childbirth.api.analyze_childbirth_plan", side_effect=AssertionError("must not recalculate")):
        restored = reloaded.get(f"{BASE}/{id}").json()["data"]
        assert restored["saved_at"] == first["saved_at"]
        assert restored["recommendations"] == data["recommendations"]
        assert restored["input"]["father"]["birth_date"] == "1990-01-02"
        assert restored["input"]["father"]["birth_place"]["display_name"] == "Hà Nội"
        assert restored["input"]["mother"]["birth_time"] is None
        history = reloaded.get(f"{BASE}/history").json()["data"]["items"]
        assert len(history) == 1
        assert history[0]["display_identity"] == "An / Bình"


def test_docx_download_contains_inputs_and_the_analyzed_recommendations(tmp_path):
    client, _ = _client(tmp_path)
    data = _create(client)
    id = data["consultation_id"]
    client.post(f"{BASE}/{id}/save")
    with patch("consulting.childbirth.api.analyze_childbirth_plan", side_effect=AssertionError("must not recalculate")):
        response = client.get(f"{BASE}/{id}/export/docx")
    assert response.status_code == 200
    assert "BTE_TuVanSinhCon" in response.headers["content-disposition"]
    assert "wordprocessingml.document" in response.headers["content-type"]
    document = Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "TƯ VẤN SINH CON" in text and "TƯ VẤN HÔN NHÂN" not in text
    assert "02/01/1990" in text and "01/02/1992" in text
    assert "Giờ Sửu" in text and "Không biết" in text and "Hà Nội" in text
    assert data["summary"]["headline"] in text
    assert document.core_properties.identifier == id
    for item in data["recommendations"]:
        assert str(item["year"]) in text
        assert f"{item['score']}/100" in text
        assert all(reason in text for reason in item["reasons"])


def test_pdf_download_uses_the_same_archive_and_cleans_temporary_file(tmp_path):
    client, repository = _client(tmp_path / "profiles")
    data = _create(client)
    id = data["consultation_id"]
    exports = []

    class FakePdfBackend:
        def html_to_pdf(self, content, output_path, *, title):
            assert "TƯ VẤN SINH CON" in content
            assert data["summary"]["headline"] in content
            assert "An" in title and "Bình" in title
            output_path.write_bytes(b"%PDF-1.4\n" + content.encode("utf-8") + b"0" * 2048)
            return 1

    def export(profile, fmt):
        artifact = export_childbirth_file(profile, fmt, pdf_backend=FakePdfBackend())
        exports.append(artifact.path)
        return artifact

    with patch("consulting.childbirth.api.export_childbirth_file", side_effect=export):
        response = client.get(f"{BASE}/{id}/export/pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert exports and not exports[0].exists()
    assert repository.get(id).result["recommendations"] == data["recommendations"]


def test_missing_profiles_and_failed_exports_return_errors(tmp_path):
    client, _ = _client(tmp_path)
    for path, method in (("unknown", "get"), ("unknown/save", "post"), ("unknown/export/docx", "get")):
        response = getattr(client, method)(f"{BASE}/{path}")
        assert response.status_code == 404
        assert response.json()["errors"][0]["code"] == "NOT_FOUND"
    data = _create(client)
    with patch("consulting.childbirth.api.export_childbirth_file", side_effect=RuntimeError("test")):
        response = client.get(f"{BASE}/{data['consultation_id']}/export/pdf")
    assert response.status_code == 500
    assert response.json()["data"] is None
    assert client.get(f"{BASE}/{data['consultation_id']}/export/zip").status_code == 422


def test_failed_save_leaves_profile_unsaved(tmp_path):
    client, repository = _client(tmp_path)
    data = _create(client)
    with patch.object(repository, "_write", side_effect=OSError("test write failed")):
        assert client.post(f"{BASE}/{data['consultation_id']}/save").status_code == 500
    assert repository.get(data["consultation_id"]).saved_at is None
    assert repository.history() == []
