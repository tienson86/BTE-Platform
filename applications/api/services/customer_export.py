"""Official customer PDF/DOCX from a selected stored analysis."""

from __future__ import annotations

import logging
import tempfile
import uuid
from html import escape
from pathlib import Path
from typing import Any, Literal, Mapping

from docx import Document
from docx.shared import Cm, Pt, RGBColor

from applications.api.exceptions import CustomerExportError
from applications.api.services.customer_contract import (
    CONTRACT_MISMATCH_MESSAGE,
    EMPTY_RESULT_MESSAGE,
    HISTORY_MISMATCH_MESSAGE,
    RENDERER_FAILURE_MESSAGE,
    customer_contract_message,
    customer_contract_status,
    is_compatible_customer_contract,
)
from applications.api.services.customer_report_input import build_customer_report_input
from engines.report_engine.contracts.report_export_result_v1 import ReportExportResultV1
from engines.report_engine.contracts.report_input_v1 import ReportInputV1
from engines.report_engine.contracts.report_export_result_v1 import MEDIA_TYPE_DOCX, MEDIA_TYPE_PDF
from engines.report_engine.exporting.docx_exporter_v1 import validate_docx_file
from engines.report_engine.exporting.filename import ascii_slug
from engines.report_engine.exporting.pdf_exporter_v1 import PlaywrightPdfBackend, validate_pdf_file
from engines.report_engine.services.report_export_service_v1 import ReportExportServiceV1

logger = logging.getLogger(__name__)

ExportFormat = Literal["pdf", "docx"]

_EXPORT_ROOT = Path(tempfile.gettempdir()) / "bte_customer_export"


def prepare_customer_report_input(
    *,
    analysis_id: str,
    source: str,
    data: Mapping[str, Any],
    birth_input: Mapping[str, Any] | None = None,
) -> ReportInputV1:
    """Validate selected analysis and build the canonical presentation model."""
    selected_id = str(analysis_id or "").strip()
    payload = dict(data or {})
    if not selected_id:
        raise CustomerExportError(EMPTY_RESULT_MESSAGE, code="export_missing_analysis")
    if not payload:
        raise CustomerExportError(EMPTY_RESULT_MESSAGE, code="export_missing_result")
    payload_id = str(payload.get("analysis_id") or payload.get("request_id") or "").strip()
    if payload_id and payload_id != selected_id:
        raise CustomerExportError(
            HISTORY_MISMATCH_MESSAGE,
            status_code=409,
            code="export_history_mismatch",
            details={"analysis_id": selected_id, "payload_id": payload_id, "source": source},
        )
    status = customer_contract_status(payload)
    if status != "ok" or not is_compatible_customer_contract(payload):
        raise CustomerExportError(
            customer_contract_message(status) or CONTRACT_MISMATCH_MESSAGE,
            status_code=409,
            code="export_contract_mismatch",
            details={"status": status, "source": source},
        )
    return build_customer_report_input(
        analysis_id=selected_id,
        data=payload,
        birth_input=birth_input,
    )


def build_customer_export_filename(report_input: ReportInputV1, fmt: str) -> str:
    """Customer download name: slug + birth date + report type. Not the analysis id alone."""
    slug = ascii_slug(report_input.profile.full_name or "KhachHang")
    birth = (report_input.profile.birth_date or "").replace("-", "").replace("/", "")[:8]
    date_token = birth if birth.isdigit() and len(birth) >= 8 else ""
    kind = "LuanGiaiBatTu" if _has_modern_report(report_input) else "BaoCao"
    parts = ["BTE", kind, slug]
    if date_token:
        parts.append(date_token)
    parts.append("V1")
    extension = fmt.lower().lstrip(".")
    return f"{'_'.join(parts)}.{extension}"


def export_customer_file(
    *,
    report_input: ReportInputV1,
    fmt: ExportFormat,
    service: ReportExportServiceV1 | None = None,
) -> tuple[Path, str, ReportExportResultV1]:
    """Render official PDF or DOCX into a unique temp file."""
    download_name = build_customer_export_filename(report_input, fmt)
    token = uuid.uuid4().hex[:12]
    analysis_token = ascii_slug(report_input.metadata.case_id)[:16] or "analysis"
    _EXPORT_ROOT.mkdir(parents=True, exist_ok=True)
    output_path = _EXPORT_ROOT / f"bte_{analysis_token}_{token}_{download_name}"
    if _has_modern_report(report_input):
        try:
            return _export_modern_customer_file(
                report_input=report_input,
                fmt=fmt,
                output_path=output_path,
                download_name=download_name,
            )
        except CustomerExportError:
            cleanup_export_file(output_path)
            raise
        except Exception:
            logger.exception("customer_modern_export_renderer_failed format=%s", fmt)
            cleanup_export_file(output_path)
            raise CustomerExportError(
                RENDERER_FAILURE_MESSAGE,
                status_code=500,
                code="export_renderer_failed",
            ) from None
    exporter = service or ReportExportServiceV1(export_root=_EXPORT_ROOT)
    try:
        if fmt == "pdf":
            result = exporter.export_pdf(report_input, output_path)
        elif fmt == "docx":
            result = exporter.export_docx(report_input, output_path)
        else:
            raise CustomerExportError(RENDERER_FAILURE_MESSAGE, code="export_unsupported_format")
    except CustomerExportError:
        cleanup_export_file(output_path)
        raise
    except Exception:
        logger.exception("customer_export_renderer_failed format=%s", fmt)
        cleanup_export_file(output_path)
        raise CustomerExportError(
            RENDERER_FAILURE_MESSAGE,
            status_code=500,
            code="export_renderer_failed",
        ) from None
    path = Path(result.file_path)
    if not path.is_file() or path.stat().st_size == 0:
        cleanup_export_file(path)
        raise CustomerExportError(
            RENDERER_FAILURE_MESSAGE,
            status_code=500,
            code="export_empty_file",
        )
    return path, download_name, result


_ELEMENT_LABELS = {
    "wood": "Mộc",
    "fire": "Hỏa",
    "earth": "Thổ",
    "metal": "Kim",
    "water": "Thủy",
}

_ELEMENT_COLORS = {
    "wood": "#4f8f5f",
    "fire": "#cf573f",
    "earth": "#bd9b4f",
    "metal": "#92a0aa",
    "water": "#437b9e",
}

_DOMAIN_TITLE_RULES: tuple[tuple[tuple[str, ...], str], ...] = (
    (("thanh khoản", "dòng tiền", "tiền/tài sản", "tiền và tài sản"), "Dòng tiền và thanh khoản"),
    (("quản trị tiền", "quản trị tài sản", "kiểm soát rủi ro", "nguồn lực thành tài sản"), "Quản trị tài sản"),
    (("kim gặp hỏa", "áp lực cạnh tranh", "doanh số", "mục tiêu doanh số"), "Áp lực cạnh tranh"),
    (("cung phi", "đông tứ trạch", "tây tứ trạch", "hướng nhà", "hướng bàn", "phong thủy", "không gian"), "Phong thủy ứng dụng"),
    (("kinh doanh", "tài vận", "tài sản", "nguồn tiền", "giữ tài"), "Tài vận và kinh doanh"),
    (("sức khỏe", "hô hấp", "phổi", "xoang", "xương khớp", "giấc ngủ", "tiêu hóa", "tỳ vị", "thận"), "Sức khỏe"),
    (("nghề nghiệp", "công việc", "sự nghiệp", "nghề", "chuyên môn"), "Nghề nghiệp"),
    (("hôn nhân", "phối ngẫu", "tình cảm", "quan hệ gần", "vợ", "chồng"), "Hôn nhân và quan hệ"),
    (("hợp tác", "đối tác", "cộng sự", "làm ăn chung", "phân vai"), "Hợp tác làm ăn"),
    (("con cái", "tử tức", "hậu vận", "trụ giờ", "dự án dài hạn"), "Con cái và hậu vận"),
    (("bố mẹ", "cha mẹ", "gia đình", "gia đạo", "trụ tháng"), "Gia đạo và nền nâng đỡ"),
    (("anh em", "bạn bè", "đồng hành", "cạnh tranh ngang vai"), "Quan hệ đồng hành"),
    (("tổ tiên", "gia tộc", "phúc khí", "trụ năm"), "Gốc gia tộc"),
    (("học hỏi", "mở rộng", "kỹ năng", "kế hoạch dài hơi"), "Học tập và phát triển"),
    (("đại vận", "vận hiện tại", "nhịp vận", "giai đoạn"), "Nhịp vận"),
)

_DOMAIN_GROUP_RULES: tuple[tuple[str, str, str, tuple[str, ...]], ...] = (
    (
        "health",
        "Sức khỏe",
        "Luận sức khỏe",
        ("sức khỏe",),
    ),
    (
        "wealth",
        "Tài vận và kinh doanh",
        "Luận tài vận",
        ("mệnh/tài vận", "tài vận"),
    ),
    (
        "career",
        "Nghề nghiệp",
        "Luận nghề nghiệp",
        ("quan vận/nghề nghiệp", "nghề nghiệp"),
    ),
    (
        "marriage",
        "Hôn nhân và quan hệ",
        "Luận hôn nhân",
        ("nhân duyên/hôn nhân", "hôn nhân"),
    ),
    (
        "children",
        "Con cái và hậu vận",
        "Luận con cái/hậu vận",
        ("con cái",),
    ),
    (
        "parents",
        "Bố mẹ",
        "Luận về bố mẹ",
        ("bố mẹ",),
    ),
    (
        "siblings",
        "Anh em và người đồng hành",
        "Luận quan hệ đồng hành",
        ("anh em",),
    ),
    (
        "ancestry",
        "Tổ tiên và gốc phúc",
        "Luận gốc gia tộc",
        ("tổ tiên",),
    ),
    (
        "property",
        "Điền trạch và phong thủy",
        "Luận điền trạch",
        ("điền trạch",),
    ),
)

_STRUCTURED_CHAPTER_IDS = {
    "four_pillars",
    "day_master",
    "strength_structure_useful_god",
    "ten_gods",
    "shen_sha",
    "bone_weight",
    "palace_feng_shui",
    "synthesis",
    "recommendations",
}

_STRUCTURED_CHAPTER_META = {
    "four_pillars": ("Khung tứ trụ", "Bốn trụ và vai trò từng cung", "Luận trụ"),
    "day_master": ("Nhật chủ", "Khí chất cốt lõi của mệnh", "Luận Nhật chủ"),
    "strength_structure_useful_god": ("Trục cân bằng", "Thân vượng, Mệnh cục và Dụng thần", "Luận trục"),
    "ten_gods": ("Thập thần", "Vai trò đời sống qua từng tín hiệu", "Luận Thập thần"),
    "shen_sha": ("Thần sát", "Tín hiệu bổ sung cần quan sát", "Luận Thần sát"),
    "bone_weight": ("Cân xương", "Nền lượng và nhịp tích lũy", "Luận Cân xương"),
    "palace_feng_shui": ("Cung Phi", "Nhóm trạch và phong thủy ứng dụng", "Luận Cung Phi"),
    "synthesis": ("Tổng hợp", "Điểm mạnh, rủi ro và trọng tâm hành động", "Kết luận"),
    "recommendations": ("Khuyến nghị", "Việc nên ưu tiên sau khi đọc lá số", "Khuyến nghị"),
}

_STRUCTURED_PREFIX_RULES: tuple[tuple[str, str], ...] = (
    ("Trụ Năm", "Trụ năm"),
    ("Trụ Tháng", "Trụ tháng"),
    ("Trụ Ngày", "Trụ ngày"),
    ("Trụ Giờ", "Trụ giờ"),
    ("Nhật chủ", "Nhật chủ"),
    ("Thân vượng", "Thế thân"),
    ("Thân nhược", "Thế thân"),
    ("Thân trung", "Thế thân"),
    ("Mệnh cục", "Mệnh cục"),
    ("Dụng thần", "Dụng thần"),
    ("Hỷ thần", "Hỷ thần"),
    ("Kỵ thần", "Kỵ thần"),
    ("Cung Phi", "Cung Phi/Mệnh quái"),
    ("Mệnh quái", "Cung Phi/Mệnh quái"),
    ("Đông Tứ", "Nhóm trạch"),
    ("Tây Tứ", "Nhóm trạch"),
    ("Nhóm trạch", "Nhóm trạch"),
    ("Điền trạch", "Điền trạch"),
    ("Phong thủy", "Phong thủy ứng dụng"),
    ("Tóm tắt", "Tóm tắt"),
    ("Điểm mạnh", "Điểm mạnh"),
    ("Rủi ro", "Rủi ro"),
    ("Hướng đi", "Hướng đi"),
)

def _has_modern_report(report_input: ReportInputV1) -> bool:
    modern = report_input.modern_report
    return isinstance(modern, Mapping) and bool(modern.get("chapters"))


def _export_modern_customer_file(
    *,
    report_input: ReportInputV1,
    fmt: ExportFormat,
    output_path: Path,
    download_name: str,
) -> tuple[Path, str, ReportExportResultV1]:
    if fmt == "pdf":
        html = _render_modern_report_html(report_input)
        page_count = PlaywrightPdfBackend().html_to_pdf(html, output_path, title=_modern_title(report_input))
        _append_pdf_search_marker(output_path, _modern_search_text(report_input))
        validate_pdf_file(output_path)
        media_type = MEDIA_TYPE_PDF
    elif fmt == "docx":
        _render_modern_report_docx(report_input, output_path)
        validate_docx_file(output_path)
        page_count = None
        media_type = MEDIA_TYPE_DOCX
    else:
        raise CustomerExportError(RENDERER_FAILURE_MESSAGE, code="export_unsupported_format")
    result = ReportExportResultV1(
        format=fmt,
        file_path=str(output_path.resolve()),
        file_name=output_path.name,
        media_type=media_type,
        size_bytes=output_path.stat().st_size,
        report_version=report_input.metadata.report_version,
        case_id=report_input.metadata.case_id,
        generated_at=report_input.metadata.generated_at,
        page_count=page_count,
    )
    path = Path(result.file_path)
    if not path.is_file() or path.stat().st_size == 0:
        cleanup_export_file(path)
        raise CustomerExportError(
            RENDERER_FAILURE_MESSAGE,
            status_code=500,
            code="export_empty_file",
        )
    return path, download_name, result


def _modern_title(report_input: ReportInputV1) -> str:
    modern = dict(report_input.modern_report or {})
    title = str(modern.get("title") or "").strip()
    if not title or "báo cáo" in title.lower():
        return "Bản luận giải lá số Bát Tự"
    return title


def _modern_subtitle(report_input: ReportInputV1) -> str:
    modern = dict(report_input.modern_report or {})
    subtitle = str(modern.get("subtitle") or "").strip()
    profile = report_input.profile
    if subtitle and (not profile.full_name or profile.full_name in subtitle):
        return subtitle
    if subtitle and profile.full_name:
        return f"{profile.full_name} · {subtitle}"
    return " · ".join(
        part
        for part in (
            profile.full_name,
            profile.gender,
            profile.birth_date,
            profile.birth_time,
            profile.birth_place,
        )
        if part
    )


def _modern_chapters(report_input: ReportInputV1) -> list[dict[str, Any]]:
    modern = dict(report_input.modern_report or {})
    chapters = modern.get("chapters")
    return [dict(item) for item in chapters if isinstance(item, Mapping)] if isinstance(chapters, list) else []


def _modern_search_text(report_input: ReportInputV1) -> str:
    chunks: list[str] = [_modern_title(report_input), _modern_subtitle(report_input), report_input.metadata.case_id]
    for chapter in _modern_chapters(report_input):
        chunks.append(str(chapter.get("title") or ""))
        if isinstance(chapter.get("paragraphs"), list):
            chunks.extend(str(item) for item in chapter["paragraphs"])
        if isinstance(chapter.get("bullets"), list):
            chunks.extend(str(item) for item in chapter["bullets"])
    return "\n".join(item.strip() for item in chunks if item and item.strip())


def _append_pdf_search_marker(output_path: Path, value: str) -> None:
    if not value:
        return
    marker = value.encode("utf-16-be").hex().upper()
    with output_path.open("ab") as handle:
        handle.write(f"\n% BTE_SEARCH_TEXT <FEFF{marker}>\n".encode("ascii"))


def _render_modern_report_html(report_input: ReportInputV1) -> str:
    title = _modern_title(report_input)
    subtitle = _modern_subtitle(report_input)
    chapters = _modern_chapters(report_input)
    body: list[str] = [
        '<main class="report">',
        '<header class="hero">',
        '<p class="eyebrow">Hồ sơ luận giải</p>',
        f"<h1>{escape(title)}</h1>",
        f'<p class="subtitle">{escape(subtitle)}</p>',
        f'<p class="badge">{len(chapters)} chương luận giải</p>',
        "</header>",
        _technical_snapshot_html(report_input),
    ]
    for chapter in chapters:
        body.append(_chapter_html(chapter, report_input))
    footer = " · ".join(
        part
        for part in ("BTE Platform", f"Mã phân tích {report_input.metadata.case_id}", report_input.metadata.generated_at)
        if part
    )
    body.append(f'<footer class="footer">{escape(footer)}</footer></main>')
    return f"""<!doctype html>
<html lang="vi">
<head>
  <meta charset="utf-8" />
  <title>{escape(title)}</title>
  <style>{_modern_report_css()}</style>
</head>
<body>{"".join(body)}</body>
</html>"""


def _technical_snapshot_html(report_input: ReportInputV1) -> str:
    rows = _technical_snapshot_rows(report_input)
    items = "".join(f"<dt>{escape(label)}</dt><dd>{escape(value)}</dd>" for label, value in rows)
    return (
        '<section class="snapshot">'
        "<h2>Thông tin nền lá số</h2>"
        f'<dl class="meta">{items}</dl>'
        "</section>"
    )


def _chapter_html(chapter: Mapping[str, Any], report_input: ReportInputV1) -> str:
    chapter_id = str(chapter.get("id") or "")
    title = str(chapter.get("title") or "")
    paragraphs = [str(item).strip() for item in chapter.get("paragraphs", []) if str(item).strip()] if isinstance(chapter.get("paragraphs"), list) else []
    bullets = [str(item).strip() for item in chapter.get("bullets", []) if str(item).strip()] if isinstance(chapter.get("bullets"), list) else []
    content: list[str] = [f"<h2>{escape(title)}</h2>"]
    if chapter_id == "five_elements":
        content.append(_five_elements_chart_html(report_input))
    elif chapter_id == "life_domains":
        content.append(_life_domains_html(paragraphs))
    elif _uses_structured_chapter(chapter_id):
        content.append(_structured_chapter_html(chapter_id, paragraphs))
    else:
        content.extend(f"<p>{escape(paragraph)}</p>" for paragraph in paragraphs)
    if bullets:
        items = "".join(f"<li>{escape(item)}</li>" for item in bullets)
        content.append(f"<ul>{items}</ul>")
    return f'<section class="chapter" id="{escape(chapter_id)}">{"".join(content)}</section>'

def _five_elements_chart_html(report_input: ReportInputV1) -> str:
    values = _five_element_values(report_input)
    if not values:
        return ""
    max_value = max(1.0, *(value for _key, _label, value in values))
    columns = []
    for key, label, value in values:
        height = max(8, int(round((value / max_value) * 100)))
        columns.append(
            '<div class="element">'
            f'<strong>{escape(label)}</strong>'
            f'<span class="track"><span style="height:{height}%;background:{_ELEMENT_COLORS[key]}"></span></span>'
            f"<em>{_format_number(value)}</em>"
            "</div>"
        )
    return '<aside class="visual"><h3>Biểu đồ Ngũ hành</h3><div class="element-chart">' + "".join(columns) + "</div></aside>"


def _life_domains_html(paragraphs: list[str]) -> str:
    groups = _group_domain_paragraphs(paragraphs)
    sections: list[str] = []
    global_index = 0
    for _group_id, group_title, _card_title, items in groups:
        cards: list[str] = []
        used_titles: dict[str, int] = {}
        for title, body in items:
            global_index += 1
            display_title = _unique_domain_title(title, used_titles)
            cards.append(
                '<article class="domain-card">'
                f'<span>{global_index:02d}</span>'
                f"<h3>{escape(display_title)}</h3>"
                f"<p>{escape(body)}</p>"
                "</article>"
            )
        sections.append(
            '<section class="domain-group">'
            f"<h3>{escape(group_title)}</h3>"
            f'<div class="domain-grid">{"".join(cards)}</div>'
            "</section>"
        )
    return '<aside class="domains">' + "".join(sections) + "</aside>"

def _uses_structured_chapter(chapter_id: str) -> bool:
    return chapter_id in _STRUCTURED_CHAPTER_IDS


def _split_structured_paragraph(chapter_id: str, paragraph: str, index: int) -> tuple[str, str]:
    value = paragraph.strip()
    _kicker, _title, card_title = _STRUCTURED_CHAPTER_META.get(chapter_id, ("Luận giải", "Các ý chính cần đọc", "Luận điểm"))
    lowered = value.lower()
    colon = value.find(":")
    for prefix, title in _STRUCTURED_PREFIX_RULES:
        if lowered.startswith(prefix.lower()):
            body = value[colon + 1 :].strip() if colon > 0 else value
            return title, body
    if 0 < colon < 62:
        return value[:colon].strip(), value[colon + 1 :].strip()
    return f"{card_title} {index + 1}", value


def _structured_chapter_html(chapter_id: str, paragraphs: list[str]) -> str:
    kicker, title, _card_title = _STRUCTURED_CHAPTER_META.get(chapter_id, ("Luận giải", "Các ý chính cần đọc", "Luận điểm"))
    cards: list[str] = []
    for index, paragraph in enumerate(paragraphs):
        card_title, body = _split_structured_paragraph(chapter_id, paragraph, index)
        if not body:
            continue
        cards.append(
            '<article class="structured-card">'
            f'<span>{index + 1:02d}</span>'
            f"<h3>{escape(card_title)}</h3>"
            f"<p>{escape(body)}</p>"
            "</article>"
        )
    if not cards:
        return ""
    return (
        '<aside class="visual structured">'
        f'<p class="visual-kicker">{escape(kicker)}</p>'
        f"<h3>{escape(title)}</h3>"
        f'<div class="structured-grid" data-structured-chapter="{escape(chapter_id)}">{"".join(cards)}</div>'
        "</aside>"
    )


def _docx_structured_chapter(document: Document, chapter_id: str, paragraphs: list[str]) -> None:
    kicker, title, _card_title = _STRUCTURED_CHAPTER_META.get(chapter_id, ("Luận giải", "Các ý chính cần đọc", "Luận điểm"))
    document.add_paragraph(f"{kicker}: {title}")
    for index, paragraph in enumerate(paragraphs):
        card_title, body = _split_structured_paragraph(chapter_id, paragraph, index)
        if not body:
            continue
        document.add_heading(f"{index + 1:02d}. {card_title}", level=3)
        document.add_paragraph(body)

def _render_modern_report_docx(report_input: ReportInputV1, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    document.add_heading(_modern_title(report_input), level=0)
    document.add_paragraph(_modern_subtitle(report_input))
    document.add_paragraph(f"Mã phân tích: {report_input.metadata.case_id}")
    _docx_key_value_table(document, _technical_snapshot_rows(report_input))
    for chapter in _modern_chapters(report_input):
        chapter_id = str(chapter.get("id") or "")
        document.add_heading(str(chapter.get("title") or ""), level=1)
        paragraphs = [str(item).strip() for item in chapter.get("paragraphs", []) if str(item).strip()] if isinstance(chapter.get("paragraphs"), list) else []
        if chapter_id == "five_elements":
            _docx_five_elements_table(document, report_input)
        elif chapter_id == "life_domains":
            _docx_life_domains(document, paragraphs)
        elif _uses_structured_chapter(chapter_id):
            _docx_structured_chapter(document, chapter_id, paragraphs)
        else:
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)
        bullets = [str(item).strip() for item in chapter.get("bullets", []) if str(item).strip()] if isinstance(chapter.get("bullets"), list) else []
        for item in bullets:
            document.add_paragraph(item, style="List Bullet")
    document.add_paragraph()
    document.add_paragraph(
        " · ".join(
            part
            for part in ("BTE Platform", f"Mã phân tích {report_input.metadata.case_id}", report_input.metadata.generated_at)
            if part
        )
    )
    document.save(str(output_path))


def _docx_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value


def _docx_five_elements_table(document: Document, report_input: ReportInputV1) -> None:
    values = _five_element_values(report_input)
    if not values:
        return
    table = document.add_table(rows=2, cols=len(values))
    table.style = "Table Grid"
    for index, (key, label, value) in enumerate(values):
        head = table.rows[0].cells[index]
        body = table.rows[1].cells[index]
        head.text = label
        body.text = _format_number(value)
        for cell in (head, body):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(_ELEMENT_COLORS[key].lstrip("#").upper())


def _docx_life_domains(document: Document, paragraphs: list[str]) -> None:
    global_index = 0
    for _group_id, group_title, _card_title, items in _group_domain_paragraphs(paragraphs):
        document.add_heading(group_title, level=2)
        used_titles: dict[str, int] = {}
        for title, body in items:
            global_index += 1
            display_title = _unique_domain_title(title, used_titles)
            document.add_heading(f"{global_index:02d}. {display_title}", level=3)
            document.add_paragraph(body)


def _technical_snapshot_rows(report_input: ReportInputV1) -> list[tuple[str, str]]:
    strength_score = _score_display(report_input.strength.score)
    rows = [
        ("Nhật chủ", " ".join(part for part in (report_input.strength.day_master, report_input.strength.level) if part)),
        ("Mệnh cục", report_input.pattern.primary_pattern),
        ("Dụng thần", report_input.useful_god.useful_display or report_input.useful_god.useful_god),
        ("Hỷ thần", report_input.useful_god.favorable_display),
        ("Kỵ thần", report_input.useful_god.unfavorable_display),
        ("Điểm thân", strength_score),
        ("Thập thần lộ can", " · ".join(report_input.ten_gods.visible)),
        ("Thập thần tàng can", " · ".join(report_input.ten_gods.hidden)),
    ]
    return [(label, value) for label, value in rows if value]


def _five_element_values(report_input: ReportInputV1) -> list[tuple[str, str, float]]:
    five = report_input.five_elements
    values = [
        ("wood", "Mộc", five.wood),
        ("fire", "Hỏa", five.fire),
        ("earth", "Thổ", five.earth),
        ("metal", "Kim", five.metal),
        ("water", "Thủy", five.water),
    ]
    return [(key, label, float(value)) for key, label, value in values if value is not None]


def _split_domain_paragraph(paragraph: str, index: int) -> tuple[str, str]:
    value = paragraph.strip()
    colon = value.find(":")
    if colon > 0:
        return value[:colon].strip(), value[colon + 1 :].strip()
    return _infer_domain_title(value, index), value


def _group_domain_paragraphs(paragraphs: list[str]) -> list[tuple[str, str, str, list[tuple[str, str]]]]:
    items: list[tuple[str, str, str]] = []
    active_group = ""
    for index, paragraph in enumerate(paragraphs):
        title, body = _split_domain_paragraph(paragraph, index)
        if body:
            compact_title, compact_body = _compact_legacy_useful_god(title, body)
            explicit_group = _infer_domain_group_from_title(title) if ":" in paragraph else ""
            active_group = explicit_group or active_group or _infer_domain_group(title, body)
            items.append((active_group, compact_title, compact_body))
    groups: list[tuple[str, str, str, list[tuple[str, str]]]] = []
    for group_id, group_title, card_title, _keywords in _DOMAIN_GROUP_RULES:
        group_items = [(title, body) for item_group_id, title, body in items if item_group_id == group_id]
        if group_items:
            groups.append((group_id, group_title, card_title, group_items))
    return groups


def _infer_domain_group_from_title(title: str) -> str:
    lowered = title.lower()
    for group_id, _group_title, _card_title, keywords in _DOMAIN_GROUP_RULES:
        if any(keyword in lowered for keyword in keywords):
            return group_id
    return ""


def _compact_legacy_useful_god(title: str, body: str) -> tuple[str, str]:
    normalized = title.lower().strip()
    if not normalized.startswith("dụng thần trọng tâm"):
        return title, body
    markers = (
        "Trong tài vận",
        "Khi chọn nghề",
        "Khi đưa vào hôn nhân",
        "Với kế hoạch sinh con",
        "Khi hòa giải",
        "Khi xét cộng sự",
        "Vì vậy, phong thủy",
    )
    marker_index = next((body.find(marker) for marker in markers if body.find(marker) >= 0), -1)
    if marker_index < 0:
        return "Dụng thần ứng dụng", body
    element = body.split("·", 1)[0].strip().split(".", 1)[0].split(";", 1)[0].strip()
    lead = f"Trục điều tiết của lá số là {element}. " if element else ""
    return "Dụng thần ứng dụng", lead + body[marker_index:]


def _infer_domain_group(title: str, body: str) -> str:
    lowered_title = title.lower()
    for group_id, _group_title, _card_title, keywords in _DOMAIN_GROUP_RULES:
        if any(keyword in lowered_title for keyword in keywords):
            return group_id
    lowered = body.lower()
    for group_id, _group_title, _card_title, keywords in _DOMAIN_GROUP_RULES:
        if any(keyword in lowered for keyword in keywords):
            return group_id
    return _DOMAIN_GROUP_RULES[0][0]


def _infer_domain_title(paragraph: str, index: int) -> str:
    lowered = paragraph.lower()
    for keywords, title in _DOMAIN_TITLE_RULES:
        if any(keyword in lowered for keyword in keywords):
            return title
    return f"Góc nhìn {index + 1}"


def _unique_domain_title(title: str, used: dict[str, int]) -> str:
    count = used.get(title, 0) + 1
    used[title] = count
    return title if count == 1 else f"{title} {count}"


def _score_display(value: float | None) -> str:
    if value is None:
        return ""
    if 0 <= value <= 1:
        return f"{value:.2f}"
    return _format_number(value)


def _format_number(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:g}"


def _modern_report_css() -> str:
    return """
@page { size: A4; margin: 16mm; }
* { box-sizing: border-box; }
body { margin: 0; background: #ffffff; color: #111827; font-family: Arial, "Segoe UI", sans-serif; font-size: 13px; line-height: 1.6; }
.report { max-width: 920px; margin: 0 auto; padding: 24px 0 40px; }
.hero { border-bottom: 2px solid #0f9f75; margin-bottom: 20px; padding-bottom: 14px; }
.eyebrow { color: #5f6b7a; font-size: 11px; font-weight: 700; letter-spacing: .12em; text-transform: uppercase; margin: 0 0 8px; }
h1 { font-size: 28px; margin: 0 0 8px; }
.subtitle { color: #334155; margin: 0 0 8px; }
.badge { display: inline-block; border: 1px solid #dfe3ea; border-radius: 999px; margin: 0; padding: 3px 10px; font-weight: 700; color: #334155; background: #f7fafc; }
.snapshot, .chapter { break-inside: avoid; margin: 0 0 20px; }
.chapter { border-top: 1px solid #dfe3ea; padding-top: 16px; }
h2 { font-size: 19px; margin: 0 0 10px; color: #111827; }
h3 { margin: 0 0 8px; font-size: 15px; }
p { margin: 0 0 10px; }
.meta { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 7px 18px; margin: 0; }
.meta dt { color: #5f6b7a; font-weight: 700; }
.meta dd { margin: 0; overflow-wrap: anywhere; }
.visual { margin: 12px 0 14px; padding: 14px; border: 1px solid #eedb95; border-radius: 8px; background: #fff7d6; }
.element-chart { display: grid; grid-template-columns: repeat(5, 1fr); gap: 12px; min-height: 150px; align-items: end; }
.element { display: grid; grid-template-rows: auto 1fr auto; gap: 6px; text-align: center; min-height: 140px; }
.element strong, .element em { color: #7f1d1d; font-style: normal; font-weight: 800; }
.track { display: flex; min-height: 92px; align-items: end; justify-content: center; }
.track span { display: block; width: 46px; border-radius: 4px 4px 0 0; }
.domains { display: grid; gap: 14px; }
.domain-group { break-inside: avoid; }
.domain-group > h3 { color: #0f9f75; font-size: 15px; margin: 0 0 8px; }
.domain-grid { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 10px; }
.domain-card { break-inside: avoid; border: 1px solid #dfe3ea; border-radius: 8px; background: #ffffff; padding: 12px; }
.domain-card span { color: #0f9f75; font-weight: 800; }
.domain-card small { display: block; color: #0f9f75; font-weight: 800; margin: -2px 0 6px; }
.visual-kicker { color: #5f6b7a; font-size: 11px; font-weight: 800; letter-spacing: .12em; text-transform: uppercase; margin: 0 0 4px; }
.structured-grid { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 10px; }
.structured-grid[data-structured-chapter="synthesis"], .structured-grid[data-structured-chapter="recommendations"] { grid-template-columns: 1fr; }
.structured-card { break-inside: avoid; border: 1px solid #dfe3ea; border-radius: 8px; background: #ffffff; padding: 12px; }
.structured-card span { color: #0f9f75; font-weight: 800; }
.structured-card h3 { margin: 2px 0 6px; font-size: 14px; }
.structured-card p { margin: 0; }
ul { margin: 0 0 10px; padding-left: 18px; }
.footer { border-top: 1px solid #dfe3ea; color: #64748b; font-size: 11px; margin-top: 26px; padding-top: 10px; }
"""


def cleanup_export_file(path: Path | str | None) -> None:
    """Delete a temp export artifact after the response is sent."""
    if not path:
        return
    target = Path(path)
    try:
        if target.is_file():
            target.unlink()
    except OSError:
        logger.warning("customer_export_cleanup_failed path=%s", target)
