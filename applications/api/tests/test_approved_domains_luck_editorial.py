from applications.api.services.approved_domains_luck_editorial import (
    life_domain_prose, luck_cycle_prose, synthesis_prose, recommendation_prose,
)
from applications.api.services.bazi_analysis_contract import LIFE_DOMAIN_KEYS, build_life_domain_sections, build_luck_cycle_narrative, build_report_chapters
from applications.api.services.orchestrator import OrchestratorService
from applications.api.services.customer_export import _render_modern_report_docx, _render_modern_report_html
from applications.api.services.customer_report_input import build_customer_report_input
from applications.api.services.result_identity import stamp_customer_result_identity
from docx import Document
from pathlib import Path


def _analyze(**changes):
    birth = dict(year=1987, month=1, day=21, hour=4, minute=31,
                 gender="male", timezone="Asia/Bangkok")
    birth.update(changes)
    return OrchestratorService().analyze(**birth)


def test_approved_son_prose_is_grounded_in_pillars_and_all_published_cycles():
    payload = _analyze()
    sections = build_life_domain_sections(payload)
    assert all(len(sections[key]["paragraphs"]) == (8 if key == "wealth" else 6 if key == "career" else 3) for key in LIFE_DOMAIN_KEYS)
    wealth = " ".join(sections["wealth"]["paragraphs"])
    assert "Tiền đến từ đâu?" in wealth and "Giáp tàng ở trụ năm" in wealth
    assert "Giáp tàng ở trụ giờ" in wealth and "Không thấy Chính Tài" in wealth
    assert "Can Ất của vận đang đi mang vai trò Chính Tài" in wealth
    assert "khách cũ giới thiệu khách mới khi họ thực sự hài lòng" in wealth
    assert "Kiếp Tài hiện rõ ở trụ tháng" in wealth
    assert "Thất Sát ở trụ năm" in sections["career"]["paragraphs"][1]
    assert "chuyên gia tư vấn" in sections["career"]["paragraphs"][3]
    assert "đào tạo, dịch vụ tư vấn" in sections["career"]["paragraphs"][4]
    luck = build_luck_cycle_narrative(payload)
    prose = luck_cycle_prose(payload, luck)
    assert len(prose) == 12  # introduction, ten cycles and one closing note
    assert any("Vận đang đi qua Ất Tỵ (2022–2031" in paragraph for paragraph in prose)
    assert any("Bính Ngọ" in paragraph and "Thất Sát" in paragraph and "không đồng nhất hai can" in paragraph for paragraph in prose)
    chapters = build_report_chapters(payload, {"life_domains": sections, "luck_cycles": luck})
    assert any("Đại vận Đinh Mùi" in paragraph and "chính là can Dụng thần" in paragraph
               for chapter in chapters if chapter["id"] == "luck_cycles" for paragraph in chapter["paragraphs"])


def test_at_ty_and_binh_ngo_have_distinct_roles_and_actions_in_every_export(tmp_path):
    payload = stamp_customer_result_identity(_analyze(), "case-distinct-cycles")
    report = build_customer_report_input(
        analysis_id="case-distinct-cycles", data=payload,
        birth_input={"year": 1987, "month": 1, "day": 21, "hour": 4, "minute": 31},
    )
    luck_chapter = next(chapter for chapter in report.modern_report["chapters"] if chapter["id"] == "luck_cycles")
    at_ty = next(line for line in luck_chapter["paragraphs"] if line.startswith("Vận đang đi qua Ất Tỵ"))
    binh_ngo = next(line for line in luck_chapter["paragraphs"] if line.startswith("Đại vận Bính Ngọ"))
    assert "Ất là Chính Tài" in at_ty and "giá bán, chi phí" in at_ty
    assert "Bính là Thất Sát" in binh_ngo and "quyền quyết định" in binh_ngo
    assert "Đinh là can Dụng thần" in binh_ngo
    assert "quyền quyết định" not in at_ty and "giá bán, chi phí" not in binh_ngo
    html = _render_modern_report_html(report)
    target = tmp_path / "distinct_cycles.docx"
    _render_modern_report_docx(report, target)
    word = " ".join(paragraph.text for paragraph in Document(target).paragraphs)
    for marker in ("Ất là Chính Tài", "Bính là Thất Sát", "giá bán, chi phí", "quyền quyết định"):
        assert marker in html and marker in word


def test_other_chart_does_not_inherit_son_pattern_palace_or_period():
    payload = _analyze(year=1997, month=7, day=1, hour=14, minute=24, gender="female")
    domains = " ".join(" ".join(life_domain_prose(key, payload)) for key in LIFE_DOMAIN_KEYS)
    luck = " ".join(luck_cycle_prose(payload, build_luck_cycle_narrative(payload)))
    assert "Cung Phi Khôn" not in domains
    assert "Mệnh cục Chính Ấn cho bạn" not in domains or payload["pattern"].get("cach_cuc") == "Chính Ấn"
    assert "Ất Tỵ (2022–2031" not in luck
    assert payload["bazi"]["day_master"] in domains
    other_closing = " ".join(synthesis_prose(payload) + recommendation_prose(payload))
    assert "Đại vận Ất Tỵ (2022–2031)" not in other_closing
    assert "Cung Phi Khôn" not in other_closing


def test_missing_day_master_keeps_existing_fallback():
    assert life_domain_prose("career", {}) == []
    assert luck_cycle_prose({}, {"cycles": [{"stem": "Bính", "branch": "Ngọ"}]}) == []


def test_pdf_html_and_word_use_the_same_approved_chapter_text(tmp_path):
    payload = stamp_customer_result_identity(_analyze(), "case-approved-editorial")
    report = build_customer_report_input(
        analysis_id="case-approved-editorial", data=payload,
        birth_input={"year": 1987, "month": 1, "day": 21, "hour": 4, "minute": 31},
    )
    html = _render_modern_report_html(report)
    path = tmp_path / "approved.docx"
    _render_modern_report_docx(report, path)
    word = " ".join(paragraph.text for paragraph in Document(path).paragraphs)
    for fragment in ("Thất Sát ở trụ năm", "Vận đang đi qua Ất Tỵ",
                     "Trục mệnh", "Chọn nhịp trong vận Ất Tỵ", "Vị trí có thể phát huy",
                     "Nhóm nghề nên khảo sát", "Tiền đến từ đâu?", "Giáp tàng ở trụ năm"):
        assert fragment in html and fragment in word


def test_money_sources_change_with_natal_finance_and_do_not_inherit_son_claims():
    huynh = _analyze(year=1966, month=9, day=24, hour=4, minute=15)
    wealth = " ".join(life_domain_prose("wealth", huynh))
    assert "Cửa Chính Tài" in wealth and "Cửa Thiên Tài" in wealth
    assert "Không thấy Chính Tài" not in wealth
    assert "Giáp tàng ở trụ năm" not in wealth
    assert "Can Ất của vận đang đi" not in wealth


def test_absent_natal_finance_does_not_become_claim_about_actual_income():
    payload = {"bazi": {"day_master": "Canh"}, "ten_gods": {"visible": [], "hidden": []}}
    wealth = " ".join(life_domain_prose("wealth", payload))
    assert "Không thấy Thiên Tài" in wealth and "Không thấy Chính Tài" in wealth
    assert "không cấm bạn làm kinh doanh" in wealth
    assert "khách cũ giới thiệu" not in wealth


def test_other_chart_career_uses_its_own_pattern_and_useful_god():
    huynh = _analyze(year=1966, month=9, day=24, hour=4, minute=15)
    career = life_domain_prose("career", huynh)
    assert "Mệnh cục Chính Tài" in career[3]
    assert "người quản lý vận hành" in career[3]
    assert "quản lý hợp đồng" in career[4]
    assert "Dụng thần Nhâm Thủy" in career[5]
    assert "chuyên gia tư vấn" not in career[3]


def test_both_portal_renderers_use_the_exported_luck_paragraphs():
    root = Path(__file__).resolve().parents[3] / "applications" / "customer_portal"
    sources = (
        root / "src/screens/commercial_dashboard/ReportDocumentSection.tsx",
        root / "static/js/bazi_report_document_mount.js",
    )
    for source in sources:
        contents = source.read_text(encoding="utf-8")
        assert '"luck_cycles",' in contents
        assert 'chapter.id === "luck_cycles"' not in contents


def test_closing_keeps_day_master_useful_god_and_current_cycle_separate():
    payload = _analyze()
    synthesis = " ".join(synthesis_prose(payload))
    actions = " ".join(recommendation_prose(payload))
    assert "Canh thuộc Dương Kim" in synthesis
    assert "Đinh Hỏa" in synthesis and "Ất Tỵ (2022–2031)" in synthesis
    assert "Kiếp Tài ở trụ tháng" in synthesis
    assert "Cuối mỗi quý" in actions
